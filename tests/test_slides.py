import pytest

from backend.app import slides
from backend.app.utils import AppError


def make_pdf(path, text):
    """A minimal valid one-page PDF containing `text`."""
    content = f"BT /F1 18 Tf 20 100 Td ({text}) Tj ET".encode()
    objs = [b"<</Type/Catalog/Pages 2 0 R>>", b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
            b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 200]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
            b"<</Length %d>>stream\n" % len(content) + content + b"\nendstream",
            b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>"]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + b"".join(b"%010d 00000 n \n" % x for x in offsets)
    out += b"trailer<</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    path.write_bytes(bytes(out))


def test_reads_pdf(tmp_path):
    make_pdf(tmp_path / "deck.pdf", "Strong induction")
    assert "Strong induction" in slides.extract_text(tmp_path / "deck.pdf")


def test_reads_powerpoint(tmp_path):
    from pptx import Presentation
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Mathematical Induction"
    deck.save(tmp_path / "deck.pptx")
    assert "[Slide 1]\nMathematical Induction" in slides.extract_text(tmp_path / "deck.pptx")


def test_broken_file_gives_a_friendly_error(tmp_path):
    (tmp_path / "broken.pdf").write_bytes(b"not a pdf")
    with pytest.raises(AppError, match="recording only"):
        slides.extract_text(tmp_path / "broken.pdf")
