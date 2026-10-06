"""Read the text out of documents (PDF, Word, PowerPoint, plain text) so they can be summarised
like a lecture, or used as slides to help the notes for a recording."""

import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from .utils import AppError

SUFFIXES = {".pdf", ".docx", ".pptx", ".txt", ".md"}
MAX_CHARS = 300_000  # ~100 pages; longer documents are cut off with a warning
WORD_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def is_document(path: Path | str) -> bool:
    return Path(path).suffix.lower() in SUFFIXES


def read(path: Path) -> list[str]:
    """The text of each page/slide (one item for Word and text files). Raises AppError if unreadable."""
    suffix = path.suffix.lower()
    if suffix not in SUFFIXES:
        raise AppError("That file type can't be read. Use a PDF, Word (.docx), PowerPoint (.pptx) or text file.")
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader
            return [page.extract_text() or "" for page in PdfReader(str(path)).pages[:500]]
        if suffix == ".pptx":
            from pptx import Presentation
            pages = []
            for slide in Presentation(str(path)).slides:
                texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
                pages.append("\n".join(t for t in texts if t.strip()))
            return pages
        if suffix == ".docx":
            with zipfile.ZipFile(path) as z:
                root = ElementTree.fromstring(z.read("word/document.xml"))
            paragraphs = ["".join(t.text or "" for t in p.iter(f"{WORD_NS}t")) for p in root.iter(f"{WORD_NS}p")]
            return ["\n".join(p for p in paragraphs if p.strip())]
        return [path.read_bytes().decode("utf-8", errors="replace")]
    except Exception:
        raise AppError("Couldn't open that document. It may be damaged or password-protected.")


def full_text(path: Path, warn=lambda _: None) -> str:
    """The whole document as one text, for making notes from it."""
    pages = [p.strip() for p in read(path) if p.strip()]
    text = re.sub(r"\n{3,}", "\n\n", "\n\n".join(pages))
    if not text.strip():
        raise AppError("The document has no readable text. If it's a scan or made of images, "
                       "try a version with selectable text.")
    if len(text) > MAX_CHARS:
        warn(f"The document is very long, so only the first ~{MAX_CHARS // 3000} pages were used.")
        text = text[:MAX_CHARS]
    return text
