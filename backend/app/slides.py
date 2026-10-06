"""Pull the text out of lecture slides (PDF or PowerPoint) to give the notes AI more context."""

from pathlib import Path

from .utils import AppError

MAX_CHARS = 40_000
SUFFIXES = {".pdf", ".pptx"}


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader
            pages = [page.extract_text() or "" for page in PdfReader(str(path)).pages[:200]]
        elif suffix == ".pptx":
            from pptx import Presentation
            pages = []
            for slide in Presentation(str(path)).slides:
                texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
                pages.append("\n".join(t for t in texts if t.strip()))
        else:
            raise AppError("Slides must be a PDF or PowerPoint (.pptx) file.")
    except AppError:
        raise
    except Exception:
        raise AppError("Couldn't read the slides file, so the notes were made from the recording only.")

    text = "\n\n".join(f"[Slide {i}]\n{t.strip()}" for i, t in enumerate(pages, 1) if t.strip())
    if not text:
        raise AppError("The slides had no readable text (maybe they're images), so the notes "
                       "were made from the recording only.")
    return text[:MAX_CHARS]
